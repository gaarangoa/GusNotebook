"""Small, read-only table snapshots. No notebook kernel or pandas required."""

import csv
import io
import json
import struct
from itertools import islice
from pathlib import Path
from zipfile import ZipFile

from .textfile import MAX_BYTES

TEXT_SUFFIXES = {".csv", ".tsv", ".json", ".jsonl", ".ndjson"}
BINARY_SUFFIXES = {".parquet", ".feather", ".xlsx"}
SUFFIXES = TEXT_SUFFIXES | BINARY_SUFFIXES
MAX_ROWS = 5000
MAX_COLUMNS = 100
MAX_CELLS = 100000
MAX_CELL_CHARS = 512
MAX_PREVIEW_CHARS = 2 * 1024 * 1024
MAX_DECODED_BYTES = 32 * 1024 * 1024
csv.field_size_limit(MAX_BYTES)


class Snapshot:
    def __init__(self, columns=None, types=None):
        self.columns = columns[:MAX_COLUMNS] if columns is not None else None
        self.types = (types or [])[:MAX_COLUMNS]
        self.rows = []
        self.notices = []
        self.truncated = False
        self.characters = self.cells = 0
        if columns is not None and len(columns) > MAX_COLUMNS:
            self.limit("Showing the first 100 columns.")

    def limit(self, message):
        self.truncated = True
        if message not in self.notices:
            self.notices.append(message)

    def cell(self, value):
        if value is None:
            return ""
        if isinstance(value, (dict, list, tuple)):
            value = json.dumps(value, ensure_ascii=False, default=str)
        elif isinstance(value, bytes):
            value = repr(value)
        else:
            value = str(value)
        if len(value) > MAX_CELL_CHARS:
            self.limit("Long cell values are shortened in the preview.")
            value = value[:MAX_CELL_CHARS] + "…"
        return value

    def add(self, values):
        if len(self.rows) >= MAX_ROWS:
            self.limit("Showing the first 5,000 rows. Search and sorting apply to this snapshot.")
            return False
        values = list(islice(iter(values), MAX_COLUMNS + 1))
        if len(values) > MAX_COLUMNS:
            self.limit("Showing the first 100 columns.")
        values = [self.cell(value) for value in values[:MAX_COLUMNS]]
        characters = sum(map(len, values))
        if self.cells + len(values) > MAX_CELLS or self.characters + characters > MAX_PREVIEW_CHARS:
            self.limit("Preview size limit reached. Search and sorting apply to this snapshot.")
            return False
        self.rows.append(values)
        self.cells += len(values)
        self.characters += characters
        return True

    def data(self, **extra):
        columns = [self.cell(value) for value in self.columns] if self.columns is not None else None
        return dict(columns=columns, types=self.types, rows=self.rows,
                    notices=self.notices, truncated=self.truncated, **extra)


def _delimiter(text, suffix, chosen):
    if chosen not in (None, "auto", ",", "\t", ";", "|"):
        raise ValueError("Unsupported table delimiter")
    if chosen not in (None, "auto"):
        return chosen
    if suffix == ".tsv":
        return "\t"
    try:
        return csv.Sniffer().sniff(text[:8192], delimiters=",\t;|").delimiter
    except csv.Error:
        return ","


def preview_text(text, suffix, delimiter=None):
    if not isinstance(text, str) or len(text.encode("utf-8")) > MAX_BYTES:
        raise ValueError("File is too large to preview (maximum 2 MB)")
    text = text.removeprefix("\ufeff")
    snapshot = Snapshot()
    if suffix in {".csv", ".tsv"}:
        chosen = _delimiter(text, suffix, delimiter)
        # Quoted delimiters, escaped quotes and multiline fields are handled by
        # the standard CSV reader; a trailing newline never adds a blank row.
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=chosen, strict=True):
            if row and not snapshot.add(row):
                break
        return snapshot.data(delimiter=chosen)
    if suffix == ".json":
        try:
            records = json.loads(text)
        except RecursionError as error:
            raise ValueError("JSON is too deeply nested to preview. Choose Source to view it.") from error
        if isinstance(records, dict):
            records = [records]
        if not isinstance(records, list):
            raise ValueError("Table preview needs a JSON object or array. Choose Source to view this value.")
        total = len(records)
        records = records[:MAX_ROWS + 1]
    elif suffix in {".jsonl", ".ndjson"}:
        records = []
        for number, line in enumerate(io.StringIO(text), 1):
            if not line.strip():
                continue
            try:
                records.append(json.loads(line))
            except (ValueError, RecursionError) as error:
                raise ValueError(f"Invalid JSON on line {number}: {error}") from error
            if len(records) > MAX_ROWS:
                break
        total = None
    else:
        raise ValueError("Unsupported table format")
    if not records:
        return Snapshot([]).data(total_rows=total)
    if all(isinstance(record, dict) for record in records):
        columns = list(dict.fromkeys(key for record in records for key in record))
        snapshot = Snapshot(columns)
        values = ([record.get(column) for column in snapshot.columns] for record in records)
    elif all(isinstance(record, list) for record in records):
        width = max(map(len, records))
        snapshot = Snapshot([f"Column {i + 1}" for i in range(min(width, MAX_COLUMNS + 1))])
        values = iter(records)
    else:
        snapshot = Snapshot(["Value"])
        values = ([record] for record in records)
    for row in values:
        if not snapshot.add(row):
            break
    return snapshot.data(total_rows=total)


def _xlsx(path, sheet):
    from openpyxl import load_workbook
    # Compressed workbooks can be much larger than their on-disk size. Refuse
    # those before openpyxl reads shared strings and styles into memory.
    with ZipFile(path) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > MAX_DECODED_BYTES:
            raise ValueError("Workbook expands beyond 32 MB; preview skipped")
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheets = workbook.sheetnames
        if not sheets:
            raise ValueError("Workbook has no worksheets")
        selected = sheet if sheet is not None else sheets[0]
        if selected not in sheets:
            raise ValueError("No such worksheet")
        worksheet = workbook[selected]
        # Ignore cached dimensions: some writers save an incorrect A1:A1.
        worksheet.reset_dimensions()
        snapshot = Snapshot()
        for row in islice(worksheet.iter_rows(max_col=MAX_COLUMNS + 1, values_only=True), MAX_ROWS + 1):
            values = list(row)
            while values and values[-1] is None:
                values.pop()
            if not snapshot.add(values):
                break
        return snapshot.data(sheets=sheets, sheet=selected,
                             note="Formula cells show their last saved result; formulas are not recalculated.")
    finally:
        workbook.close()


def _parquet(path):
    import pyarrow.parquet as parquet
    with parquet.ParquetFile(path) as reader:
        metadata = reader.metadata
        expanded = sum(metadata.row_group(i).total_byte_size for i in range(metadata.num_row_groups))
        if expanded > MAX_DECODED_BYTES:
            raise ValueError("Parquet data expands beyond 32 MB; preview skipped")
        schema = reader.schema_arrow
        snapshot = Snapshot(schema.names, [str(field.type) for field in schema])
        for batch in reader.iter_batches(batch_size=256, columns=snapshot.columns,
                                         use_threads=False):
            columns = [column.to_pylist() for column in batch.columns]
            for row in zip(*columns):
                if not snapshot.add(row):
                    return snapshot.data(total_rows=metadata.num_rows)
        return snapshot.data(total_rows=metadata.num_rows)


def _flatbuffer_field(buffer, table, slot):
    """Locate a field without decoding Arrow's data buffers.

    Layouts: apache/arrow format/File.fbs and format/Message.fbs. This only
    inspects the IPC size metadata; pyarrow still handles all data decoding.
    """
    vtable = table - struct.unpack_from("<i", buffer, table)[0]
    size = struct.unpack_from("<H", buffer, vtable)[0]
    entry = 4 + slot * 2
    offset = struct.unpack_from("<H", buffer, vtable + entry)[0] if entry < size else 0
    return table + offset if offset else None


def _flatbuffer_target(buffer, field):
    return field + struct.unpack_from("<I", buffer, field)[0]


def _check_feather_size(source, arrow):
    """Refuse compressed IPC files before allocating their expanded buffers."""
    size = source.size()
    footer_size = struct.unpack("<I", source.read_at(4, size - 10))[0]
    if not 0 < footer_size <= size - 10:
        raise ValueError("Invalid Feather footer")
    footer = source.read_at(footer_size, size - 10 - footer_size)
    root = struct.unpack_from("<I", footer)[0]
    expanded = 0
    for slot in (2, 3):  # Dictionary and record-batch block vectors.
        field = _flatbuffer_field(footer, root, slot)
        if field is None:
            continue
        vector = _flatbuffer_target(footer, field)
        count = struct.unpack_from("<I", footer, vector)[0]
        for index in range(count):
            offset = struct.unpack_from("<q", footer, vector + 4 + index * 24)[0]
            source.seek(offset)
            message = arrow.ipc.read_message(source)
            metadata = message.metadata
            header_field = _flatbuffer_field(metadata, struct.unpack_from("<I", metadata)[0], 2)
            header = _flatbuffer_target(metadata, header_field)
            if message.type == "dictionary":
                header = _flatbuffer_target(metadata, _flatbuffer_field(metadata, header, 1))
            buffers_field = _flatbuffer_field(metadata, header, 2)
            if buffers_field is None:
                continue
            buffers = _flatbuffer_target(metadata, buffers_field)
            compressed = _flatbuffer_field(metadata, header, 3) is not None
            for position in range(struct.unpack_from("<I", metadata, buffers)[0]):
                start, length = struct.unpack_from("<qq", metadata, buffers + 4 + position * 16)
                if compressed and length:
                    decoded = struct.unpack_from("<q", message.body, start)[0]
                    length = length - 8 if decoded == -1 else decoded
                if length < 0:
                    raise ValueError("Invalid Feather buffer size")
                expanded += length
                if expanded > MAX_DECODED_BYTES:
                    raise ValueError("Feather data expands beyond 32 MB; preview skipped")
    source.seek(0)


def _feather(path):
    import pyarrow as arrow
    import pyarrow.feather as feather
    with arrow.memory_map(str(path), "r") as source:
        if source.read(6) == b"ARROW1":
            _check_feather_size(source, arrow)
            source.seek(0)
            # Feather v2 is an Arrow IPC file; read individual batches rather
            # than feather.read_table(), which combines the entire dataset.
            schema = arrow.ipc.open_file(source).schema
            options = arrow.ipc.IpcReadOptions(included_fields=list(range(min(len(schema), MAX_COLUMNS))),
                                              use_threads=False)
            reader = arrow.ipc.open_file(source, options=options)
            snapshot = Snapshot(schema.names, [str(field.type) for field in schema])
            for index in range(reader.num_record_batches):
                batch = reader.get_batch(index)
                if batch.nbytes > MAX_DECODED_BYTES:
                    raise ValueError("Feather batch expands beyond 32 MB; preview skipped")
                # Only convert the preview rows to Python objects.
                columns = [column.to_pylist() for column in batch.slice(0, MAX_ROWS + 1).columns]
                for row in zip(*columns):
                    if not snapshot.add(row):
                        return snapshot.data()
            return snapshot.data()
        source.seek(0)
        # Feather v1 is uncompressed, bounded by the same 2 MB file limit.
        table = feather.read_table(source, use_threads=False)
        snapshot = Snapshot(table.column_names, [str(field.type) for field in table.schema])
        columns = [column.to_pylist() for column in table.slice(0, MAX_ROWS + 1).columns[:MAX_COLUMNS]]
        for row in zip(*columns):
            if not snapshot.add(row):
                break
        return snapshot.data(total_rows=table.num_rows)


def preview_file(path, sheet=None):
    path = Path(path)
    if path.suffix.lower() not in BINARY_SUFFIXES:
        raise ValueError("Unsupported table format")
    if not path.is_file():
        raise ValueError("No such table file")
    if path.stat().st_size > MAX_BYTES:
        raise ValueError(f"{path.name} is too large to preview (maximum 2 MB); file was not loaded")
    try:
        if path.suffix.lower() == ".xlsx":
            return _xlsx(path, sheet)
        if path.suffix.lower() == ".parquet":
            return _parquet(path)
        return _feather(path)
    except ImportError as error:
        raise ValueError("Table reader is unavailable. Update the server's GusNotebook installation.") from error
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(f"Could not preview {path.name}: {error}") from error
