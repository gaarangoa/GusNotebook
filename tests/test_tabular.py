"""Parsing, bounded decoding, and authenticated table-preview routes."""

import csv
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
import pyarrow as arrow
import pyarrow.feather as feather
import pyarrow.parquet as parquet

from gusnotebook import tabular
from gusnotebook.app import create_app, close_app
from gusnotebook.textfile import kind_of


class TableTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name).resolve()

    def tearDown(self):
        self.temporary.cleanup()

    def test_csv_quotes_newlines_bom_and_delimiters(self):
        text = '\ufeffname,value,note\r\n"a,b",42,"first\nsecond"\r\n"quote ""here""",,last\r\n'
        result = tabular.preview_text(text, '.csv')
        self.assertEqual(result['rows'], [
            ['name', 'value', 'note'], ['a,b', '42', 'first\nsecond'], ['quote "here"', '', 'last']])
        self.assertFalse(result['truncated'])
        for separator in ('\t', ';', '|'):
            text = f'name{separator}value\nA{separator}7\n'
            self.assertEqual(tabular.preview_text(text, '.csv')['delimiter'], separator)
        self.assertEqual(tabular.preview_text('a\tb\nx\ty', '.tsv')['rows'], [['a', 'b'], ['x', 'y']])
        self.assertEqual(tabular.preview_text('a|b\nc|d', '.csv', '|')['rows'][1], ['c', 'd'])
        with self.assertRaises(ValueError):
            tabular.preview_text('a,b', '.csv', 'bad')
        with self.assertRaises(csv.Error):
            tabular.preview_text('a,b\n"unclosed', '.csv')

    def test_json_records_nested_values_and_json_lines(self):
        text = json.dumps([{'name': 'A', 'nested': {'x': 1}}, {'value': 2}])
        result = tabular.preview_text(text, '.json')
        self.assertEqual(result['columns'], ['name', 'nested', 'value'])
        self.assertEqual(result['rows'], [['A', '{"x": 1}', ''], ['', '', '2']])
        self.assertEqual(tabular.preview_text('{"a":1}', '.json')['rows'], [['1']])
        self.assertEqual(tabular.preview_text('[[1,2],[3]]', '.json')['columns'], ['Column 1', 'Column 2'])
        self.assertEqual(tabular.preview_text('[1,"two",null]', '.json')['rows'], [['1'], ['two'], ['']])
        self.assertEqual(tabular.preview_text('[]', '.json')['rows'], [])
        for suffix in ('.jsonl', '.ndjson'):
            result = tabular.preview_text('{"a":1}\n\n{"a":2}\n', suffix)
            self.assertEqual(result['rows'], [['1'], ['2']])
            with self.assertRaisesRegex(ValueError, 'line 2'):
                tabular.preview_text('{"a":1}\nbad', suffix)
        with self.assertRaisesRegex(ValueError, 'object or array'):
            tabular.preview_text('null', '.json')

    def test_snapshot_limits_rows_columns_cells_and_characters(self):
        result = tabular.preview_text('value\n' + '1\n' * 6000, '.csv')
        self.assertEqual(len(result['rows']), 5000)
        self.assertTrue(result['truncated'])
        result = tabular.preview_text(','.join(str(i) for i in range(101)), '.csv')
        self.assertEqual(len(result['rows'][0]), 100)
        self.assertTrue(result['truncated'])
        result = tabular.preview_text('x\n' + 'a' * 150000, '.csv')
        self.assertEqual(len(result['rows'][1][0]), 513)
        self.assertTrue(result['truncated'])
        with patch.object(tabular, 'MAX_CELLS', 4):
            result = tabular.preview_text('a,b\n1,2\n3,4', '.csv')
            self.assertEqual(len(result['rows']), 2)
            self.assertTrue(result['truncated'])
        with patch.object(tabular, 'MAX_PREVIEW_CHARS', 4):
            result = tabular.preview_text('ab\ncd\nef', '.csv')
            self.assertEqual(result['rows'], [['ab'], ['cd']])
            self.assertTrue(result['truncated'])
        with self.assertRaisesRegex(ValueError, 'maximum 2 MB'):
            tabular.preview_text('x' * (tabular.MAX_BYTES + 1), '.csv')

    def test_parquet_and_both_feather_versions(self):
        table = arrow.table({'name': ['A', 'B'], 'value': [10, 2], 'missing': [None, 'present']})
        path = self.root / 'data.parquet'
        parquet.write_table(table, path)
        expected = [['A', '10', ''], ['B', '2', 'present']]
        self.assertEqual(tabular.preview_file(path)['rows'], expected)
        for version, compression in ((1, None), (2, 'uncompressed'), (2, 'lz4'), (2, 'zstd')):
            path = self.root / 'data.feather'
            feather.write_feather(table, path, version=version, compression=compression)
            result = tabular.preview_file(path)
            self.assertEqual(result['rows'], expected)
            self.assertEqual(result['types'], ['string', 'int64', 'string'])

    def test_feather_dictionary_and_multiple_batches(self):
        path = self.root / 'dictionary.feather'
        table = arrow.table({'category': arrow.array(['a', 'b'] * 3000).dictionary_encode()})
        feather.write_feather(table, path, compression='zstd', chunksize=1000)
        result = tabular.preview_file(path)
        self.assertEqual(len(result['rows']), 5000)
        self.assertEqual(result['rows'][:2], [['a'], ['b']])
        self.assertTrue(result['truncated'])

    def test_binary_wide_and_long_data_is_bounded(self):
        table = arrow.table({f'c{i}': ['x' * 1024] for i in range(101)})
        for suffix in ('parquet', 'feather'):
            path = self.root / ('wide.' + suffix)
            (parquet.write_table if suffix == 'parquet' else feather.write_feather)(table, path)
            result = tabular.preview_file(path)
            self.assertEqual(len(result['columns']), 100)
            self.assertEqual(len(result['rows'][0]), 100)
            self.assertEqual(len(result['rows'][0][0]), 513)
            self.assertTrue(result['truncated'])

    def test_compressed_data_is_refused_before_decoding(self):
        table = arrow.table({'value': ['x' * 4096] * 50})
        pq_path = self.root / 'compressed.parquet'
        feather_path = self.root / 'compressed.feather'
        parquet.write_table(table, pq_path)
        feather.write_feather(table, feather_path, compression='zstd')
        with patch.object(tabular, 'MAX_DECODED_BYTES', 1024):
            for path in (pq_path, feather_path):
                with self.assertRaisesRegex(ValueError, 'expands beyond'):
                    tabular.preview_file(path)
        workbook = self.root / 'compressed.xlsx'
        with ZipFile(workbook, 'w', ZIP_DEFLATED) as archive:
            archive.writestr('xl/sharedStrings.xml', 'x' * 2048)
        with patch.object(tabular, 'MAX_DECODED_BYTES', 1024), patch('openpyxl.load_workbook') as load:
            with self.assertRaisesRegex(ValueError, 'expands beyond'):
                tabular.preview_file(workbook)
            load.assert_not_called()

    def test_xlsx_sheets_and_incorrect_dimensions(self):
        path = self.root / 'book.xlsx'
        workbook = Workbook()
        workbook.active.title = 'Results'
        workbook.active.append(['name', 'value'])
        workbook.active.append(['A', 42])
        second = workbook.create_sheet('Other # α')
        second.append(['Date', 'Formula'])
        second.append(['today', '=1+1'])
        workbook.save(path)
        first = tabular.preview_file(path)
        self.assertEqual(first['sheets'], ['Results', 'Other # α'])
        self.assertEqual(first['rows'], [['name', 'value'], ['A', '42']])
        self.assertEqual(tabular.preview_file(path, 'Other # α')['rows'][1], ['today'])
        with self.assertRaisesRegex(ValueError, 'No such worksheet'):
            tabular.preview_file(path, 'Missing')
        # Some generators incorrectly advertise only A1 while writing more rows.
        with ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries['xl/worksheets/sheet1.xml'] = entries['xl/worksheets/sheet1.xml'].replace(b'ref="A1:B2"', b'ref="A1:A1"')
        with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
            for name, data in entries.items():
                archive.writestr(name, data)
        self.assertEqual(tabular.preview_file(path)['rows'], first['rows'])

    def test_oversize_binary_skips_reader_and_corrupt_files_fail_clearly(self):
        for suffix in ('parquet', 'feather', 'xlsx'):
            path = self.root / ('large.' + suffix)
            path.write_bytes(b'x' * (tabular.MAX_BYTES + 1))
            with self.assertRaisesRegex(ValueError, 'file was not loaded'):
                tabular.preview_file(path)
            path.write_bytes(b'corrupt')
            with self.assertRaises(ValueError):
                tabular.preview_file(path)


class TableRouteTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name).resolve()
        self.app = create_app({'WORK_DIR': str(self.root), 'STATE_DIR': str(self.root / 'state'),
                               'START_WATCHERS': False, 'AUTH_TOKEN': 'test'})
        self.client = self.app.test_client()
        self.client.post('/auth', headers={'Authorization': 'Bearer test'})

    def tearDown(self):
        close_app(self.app)
        self.temporary.cleanup()

    def test_open_source_preview_and_no_disk_writes(self):
        path = self.root / 'data # α.csv'
        original = 'name,value\nA,10\nB,2\n'
        path.write_text(original)
        data = self.client.post('/api/open', json={'path': str(path)}).json
        self.assertEqual(data['kind'], 'text')
        self.assertEqual(data['table_preview']['rows'][1], ['A', '10'])
        response = self.client.post('/api/table-preview', json={'path': str(path), 'text': 'name|value\nC|3', 'delimiter': '|'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['rows'][1], ['C', '3'])
        self.assertEqual(path.read_text(), original)
        self.assertEqual(self.app.extensions['gusnotebook'].kernels.status(str(path)), 'stopped')

    def test_binary_restore_refusal_and_authentication(self):
        path = self.root / 'data.PARQUET'
        parquet.write_table(arrow.table({'value': [1, 2]}), path)
        response = self.client.post('/api/open', json={'path': str(path)})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['kind'], 'table')
        self.assertIn({'path': str(path), 'kind': 'table'}, self.client.get('/api/tabs').json['tabs'])
        oversized = self.root / 'large.feather'
        oversized.write_bytes(b'x' * (tabular.MAX_BYTES + 1))
        self.assertEqual(self.client.post('/api/open', json={'path': str(oversized)}).status_code, 400)
        self.assertNotIn(str(oversized), [entry['path'] for entry in self.client.get('/api/tabs').json['tabs']])
        anonymous = self.app.test_client()
        self.assertEqual(anonymous.post('/api/table-preview', json={'path': str(path)}).status_code, 401)
        for suffix in ('csv', 'tsv', 'json', 'jsonl', 'ndjson'):
            self.assertEqual(kind_of('example.' + suffix), 'text')
        for suffix in ('parquet', 'feather', 'xlsx'):
            self.assertEqual(kind_of('example.' + suffix), 'table')

    def test_invalid_text_can_be_edited_and_bad_parameters_are_rejected(self):
        path = self.root / 'broken.json'
        path.write_text('{bad')
        opened = self.client.post('/api/open', json={'path': str(path)})
        self.assertEqual(opened.status_code, 200)
        self.assertEqual(opened.json['text'], '{bad')
        self.assertNotIn('table_preview', opened.json)
        for body in ({}, {'path': 'relative.csv'}, {'path': str(path), 'text': []},
                     {'path': str(self.root / 'book.xlsx'), 'sheet': []},
                     {'path': str(self.root / 'file.exe')}):
            response = self.client.post('/api/table-preview', json=body)
            self.assertEqual(response.status_code, 400)

    def test_json_opens_as_source_without_table_snapshot_and_is_size_limited(self):
        for suffix in ('json', 'jsonl', 'ndjson'):
            path = self.root / ('document.' + suffix)
            source = '{"nested":{"value":9007199254740993}}\n'
            path.write_text(source)
            opened = self.client.post('/api/open', json={'path': str(path)}).json
            self.assertEqual(opened['kind'], 'text')
            self.assertEqual(opened['text'], source)
            self.assertNotIn('table_preview', opened)
        large = self.root / 'large.json'
        large.write_text('"' + 'x' * tabular.MAX_BYTES + '"')
        response = self.client.post('/api/open', json={'path': str(large)})
        self.assertEqual(response.status_code, 400)
        self.assertIn('maximum 2 MB', response.json['error'])
        self.assertNotIn(str(large), self.app.extensions['gusnotebook'].texts.paths())
        self.assertNotIn(str(large), [entry['path'] for entry in self.client.get('/api/tabs').json['tabs']])


if __name__ == '__main__':
    unittest.main()
